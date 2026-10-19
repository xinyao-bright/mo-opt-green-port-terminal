"""Isolated UI/API smoke test and preview server; never opens the user database.

Run: python tools/ui_smoke.py
Preview: python tools/ui_smoke.py --serve --port 5001
Preview login: qa_viewer / Preview2026!
"""
from pathlib import Path
import argparse
import hashlib
import importlib
import io
import json
import math
import sys
import tempfile

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))


def fingerprint(path):
    return hashlib.sha256(path.read_bytes()).hexdigest() if path.exists() else None


def load_isolated_app(database):
    # app.py assigns its SQLite URL before db.init_app(), so substitute only at
    # engine initialization, before its create_all() executes. Source is intact.
    from flask_sqlalchemy import SQLAlchemy
    original_init = SQLAlchemy.init_app

    def isolated_init(self, app):
        app.config.update(
            SQLALCHEMY_DATABASE_URI='sqlite:///' + database.as_posix(),
            TESTING=True,
            TEMPLATES_AUTO_RELOAD=True,
            SECRET_KEY='local-ui-qa-only',
        )
        return original_init(self, app)

    SQLAlchemy.init_app = isolated_init
    try:
        return importlib.import_module('app')
    finally:
        SQLAlchemy.init_app = original_init


def check(response, expected=200):
    assert response.status_code == expected, (response.status_code, response.get_data(as_text=True)[:500])
    return response.get_json() if response.is_json else response.data


def seed(module, preview=False, csv_path=None):
    client = module.app.test_client()
    check(client.get('/'), 302)
    check(client.get('/login'))
    check(client.get('/register'))
    check(client.post('/register', json={'username': 'qa_viewer', 'email': 'qa-viewer@example.invalid', 'password': 'Preview2026!'}))
    check(client.post('/login', json={'username': 'qa_viewer', 'password': 'Preview2026!', 'remember': False}))
    for page in ('/', '/port-config', '/history'):
        check(client.get(page))
    assert check(client.get('/api/ships')) == []
    check(client.post('/api/schedule', json={}), 400)
    sample = (Path(csv_path) if csv_path else ROOT / 'test_ships.csv').read_bytes()
    if not preview and csv_path is None:
        sample = b'\n'.join(sample.splitlines()[:4]) + b'\n'
    check(client.post('/api/ships/import', data={'file': (io.BytesIO(sample), 'ships.csv')}, content_type='multipart/form-data'))
    ships = check(client.get('/api/ships'))
    expected_count = len(module.pd.read_csv(io.BytesIO(sample))) if csv_path else (10 if preview else 3)
    assert len(ships) == expected_count
    if csv_path:
        print(f'CSV source: {csv_path} | {len(ships)} records')
        print('First imported record: ' + json.dumps(ships[0] if ships else {}, ensure_ascii=False))
    config = {'name': '示范港区 · 标准作业配置', 'totalBerths': 10, 'totalQcs': 15,
              'qcEfficiency': 48, 'maxQcPerVessel': 3, 'safetyInterval': 0.167,
              'berthConfig': [{'id': i, 'name': f'B{i:02d}', 'length': 350} for i in range(1, 11)]}
    config_id = check(client.post('/api/port-config', json=config))['id']
    check(client.put(f'/api/port-config/{config_id}', json={'name': config['name']}))
    check(client.post(f'/api/port-config/{config_id}/activate'))
    assert check(client.get('/api/port-config'))[0]['isActive']
    check(client.delete(f'/api/port-config/{config_id}'), 400)
    module.np.random.seed(42)
    module.random.seed(42)
    result = check(client.post('/api/schedule', json={
        'w_co2': 0.4, 'w_time': 0.3, 'w_cost': 0.3,
        'pop_size': 20 if preview else 8, 'max_iter': 12 if preview else 3,
    }))
    assert result['solutions'] and result['historyId']
    assert sum(bool(s['isRecommended']) for s in result['solutions']) == 1
    assert all(len(s['assignments']) == len(ships) for s in result['solutions'])
    print(f"PASS: login, pages, CSV import, configuration editing/activation, real NSGA-II ({len(ships)} ships, {len(result['solutions'])} solutions)")
    return client, result, config_id


def smoke(module, csv_path=None):
    client, result, config_id = seed(module, csv_path=csv_path)
    for solution in result['solutions']:
        assert solution['objectiveModel'] == 'physical-v2'
        assert solution['isFeasible'] and solution['constraintPenalty'] == 0
        assert solution['accountingConfig']['auv_money'] == 10
        assert solution['objectiveParameters']['gridEmissionFactor'] == 0.5306
        totals = solution['objectiveBreakdown']
        for key in totals:
            assert math.isclose(totals[key], sum(a['objectiveBreakdown'][key] for a in solution['assignments']), abs_tol=1e-4)
        assert math.isclose(totals['co2'], sum(totals[k] for k in ('waitingCarbon', 'serviceCarbon', 'electricityCarbon')), abs_tol=1e-4)
        assert math.isclose(totals['cost'], sum(totals[k] for k in ('waitingCost', 'qcOperatingCost', 'electricityCost', 'coordinationCost')), abs_tol=1e-4)
        for key in ('co2', 'stayTime', 'cost'):
            assert math.isclose(solution[key], totals[key], abs_tol=0.0051)
        rec = solution['recommendation']
        assert math.isclose(rec['score'], math.sqrt(sum(w*z*z for w,z in zip(rec['weights'], rec['normalizedObjectives']))))
    assert min(result['solutions'], key=lambda s: s['recommendation']['score'])['isRecommended']
    for parameters in ({'gridEmissionFactor': -1}, {'unknown': 1}, []):
        check(client.post('/api/schedule', json={'objectiveParameters': parameters}), 400)
    # Exercise the no-feasible response with a deliberately injected solver result.
    original_solver = module.nsga2_solver
    module.nsga2_solver = lambda vessels, **kwargs: (module.np.zeros((1, 3)), module.np.zeros((1, 4*len(vessels))), module.np.array([50000.0]))
    try:
        failure = check(client.post('/api/schedule', json={}), 422)
        assert failure['minimumPenalty'] == 50000
    finally:
        module.nsga2_solver = original_solver
    print('PASS: v2 component totals, parameter/config snapshots, recommendation, invalid inputs, no-feasible response')
    history_id = result['historyId']
    assert len(check(client.get('/api/history'))) == 1
    detail = check(client.get(f'/api/history/{history_id}'))
    assert detail['solutions'] == result['solutions']
    solution_id = result['solutions'][0]['id']
    exported = client.get(f'/api/history/{history_id}/export?solution={solution_id}')
    check(exported)
    from openpyxl import load_workbook
    book = load_workbook(io.BytesIO(exported.data), read_only=True)
    assert book.sheetnames == ['方案总览', '逐船调度明细']
    assert book['逐船调度明细'].max_row == len(result['solutions'][0]['assignments']) + 1
    book.close()
    other = module.app.test_client()
    check(other.post('/register', json={'username': 'qa_other', 'email': 'qa-other@example.invalid', 'password': 'Preview2026!'}))
    check(other.post('/login', json={'username': 'qa_other', 'password': 'Preview2026!'}))
    assert check(other.get('/api/ships')) == []
    assert check(other.get('/api/history')) == []
    assert check(other.get('/api/port-config')) == []
    check(other.get(f'/api/history/{history_id}'), 404)
    check(other.get(f'/api/history/{history_id}/export'), 404)
    check(other.put(f'/api/port-config/{config_id}', json={'name': 'forbidden'}), 404)
    spare = check(client.post('/api/port-config', json={'name': 'Disposable QA', 'totalBerths': 1, 'totalQcs': 1, 'berthConfig': [{'id': 1, 'name': 'Q1', 'length': 350}]}))['id']
    check(client.delete(f'/api/port-config/{spare}'))
    check(client.delete(f'/api/history/{history_id}'))
    check(client.post('/api/ships/clear'))
    assert check(client.get('/api/history')) == []
    assert check(client.get('/api/ships')) == []
    check(client.get('/logout'), 302)
    check(client.get('/history'), 302)
    print('PASS: history, Excel export, user isolation, deletion, clearing, logout')


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--serve', action='store_true')
    parser.add_argument('--port', type=int, default=5001)
    parser.add_argument('--csv', type=Path, help='Import all records from this CSV into the isolated database')
    args = parser.parse_args()
    protected = [ROOT / name for name in ('app.py', 'auth.py', 'models.py', 'instance/database.db')]
    before = {path: fingerprint(path) for path in protected}
    with tempfile.TemporaryDirectory(prefix='eco_ui_qa_') as temporary:
        module = load_isolated_app(Path(temporary) / 'qa.sqlite')
        try:
            if args.serve:
                seed(module, preview=True, csv_path=args.csv)
                module.app.config['TESTING'] = False
                print(f'Isolated preview: http://127.0.0.1:{args.port} | qa_viewer / Preview2026!', flush=True)
                module.app.run(host='127.0.0.1', port=args.port, debug=False, use_reloader=False)
            else:
                smoke(module, csv_path=args.csv)
        finally:
            with module.app.app_context():
                module.db.session.remove()
                module.db.engine.dispose()
            assert all(fingerprint(path) == digest for path, digest in before.items()), 'Protected source/database changed during QA'
            print('PASS: backend source and user database fingerprints unchanged')


if __name__ == '__main__':
    main()
